#!/usr/bin/env python3
"""
refine_columns.py -- local post-processing: turn a raw OneTrust CLI export
(~200 flattened columns) into a clean CSV with exactly the 19 headers below,
in this order. No API calls; works entirely on files already on disk.

    Risk GUID | ID | Risk name | Source | Cat | Risk owners | Description |
    Treatment plan | Organization | Stage | Inherent risk score |
    Inherent risk level | Residual risk score | Residual risk level |
    Date created | Category | Result | Date closed | Risk approver

'Cat' is not in the export. It is filled by looking up each risk's Source
(supplier) name in Supplier_Category_List.xlsx and taking that supplier's
category. No match -> blank.

Usage:
    python refine_columns.py --in raw_export.csv --out refined.csv
    python refine_columns.py --in raw.csv --out out.csv \
        --cat-file Supplier_Category_List.xlsx
    # override auto-detection of the lookup sheet's columns:
    python refine_columns.py --in raw.csv --out out.csv \
        --cat-file list.xlsx --cat-key "Supplier Name" --cat-value "Category"
"""

import argparse
import csv
import re
import sys
from pathlib import Path

# ---- output shape: (exact header, [candidate raw columns, best first]) ------ #
# 'Cat' has no raw source; it is resolved from the supplier lookup instead.
HEADER_MAP = [
    ("Risk GUID",           ["id"]),
    ("ID",                  ["number"]),
    ("Risk name",           ["name"]),
    ("Source",              ["source.name"]),
    ("Cat",                 []),                      # supplier-category lookup
    ("Risk owners",         ["riskOwnersName", "riskOwner"]),
    ("Description",         ["description"]),
    ("Treatment plan",      ["recommendation", "treatment"]),
    ("Organization",        ["orgGroup.name"]),
    ("Stage",               ["stage.name"]),
    ("Inherent risk score", ["inherentRiskLevel.riskScore"]),
    ("Inherent risk level", ["inherentRiskLevel.level"]),
    ("Residual risk score", ["riskScore"]),
    ("Residual risk level", ["level"]),
    ("Date created",        ["createdUTCDateTime"]),
    ("Category",            ["riskCategoryNames", "categories"]),
    ("Result",              ["result"]),
    ("Date closed",         ["dateClosed"]),
    ("Risk approver",       ["riskApprovers"]),
]
CAT_HEADER = "Cat"
MATCH_SOURCE_HEADER = "Source"      # which output column supplies the lookup key


# --------------------------------------------------------------------------- #
def load_csv(path):
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with open(path, newline="", encoding=enc) as f:
                r = csv.DictReader(f)
                return (r.fieldnames or []), list(r)
        except UnicodeDecodeError:
            continue
    sys.exit(f"Could not decode {path} as CSV.")


def load_xlsx(path):
    try:
        from openpyxl import load_workbook          # pip install openpyxl
    except ImportError:
        sys.exit("Reading .xlsx needs openpyxl (pip install openpyxl), or save "
                 "Supplier_Category_List as CSV and pass that to --cat-file.")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(it)]
    rows = []
    for raw in it:
        rows.append({headers[i]: ("" if v is None else v)
                     for i, v in enumerate(raw) if i < len(headers)})
    return headers, rows


def load_table(path):
    p = Path(path)
    if not p.exists():
        sys.exit(f"Not found: {path}")
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        return load_xlsx(p)
    return load_csv(p)


# --------------------------------------------------------------------------- #
_ws = re.compile(r"\s+")
_paren = re.compile(r"\s*\([^)]*\)\s*$")     # trailing "(External)" / "(Inactive)"


def nkey(s):
    """Normalized supplier-name key: drop trailing parenthetical, lowercase,
    collapse whitespace."""
    s = _paren.sub("", str(s).strip())
    return _ws.sub(" ", s).strip().lower()


def detect_cat_cols(headers, rows, key_arg, val_arg):
    """Pick the (name column, category column) in the lookup sheet."""
    if key_arg and val_arg:
        return key_arg, val_arg
    lower = {h.lower(): h for h in headers}

    def find(subs):
        for h in headers:
            hl = h.lower()
            if any(s in hl for s in subs):
                return h
        return None

    key = key_arg or find(["supplier", "vendor", "name", "source"])
    val = val_arg or find(["categor", "cat", "type", "segment"])
    if not key or not val or key == val:
        # fall back to first two columns
        cols = [h for h in headers if h]
        if len(cols) >= 2:
            key = key_arg or cols[0]
            val = val_arg or (cols[1] if cols[1] != key else cols[0])
    if not key or not val:
        sys.exit("Could not identify name/category columns in the lookup sheet. "
                 "Pass --cat-key and --cat-value explicitly.")
    return key, val


def build_cat_lookup(path, key_arg, val_arg):
    headers, rows = load_table(path)
    key, val = detect_cat_cols(headers, rows, key_arg, val_arg)
    lut = {}
    for r in rows:
        k = nkey(r.get(key, ""))
        v = str(r.get(val, "")).strip()
        if k and v and k not in lut:
            lut[k] = v
    print(f"  supplier lookup: key='{key}' value='{val}' -> {len(lut)} suppliers")
    return lut


# --------------------------------------------------------------------------- #
def resolve_sources(raw_headers):
    """For each output header, choose the raw column present in the input."""
    present = set(raw_headers)
    chosen, missing = [], []
    for header, cands in HEADER_MAP:
        col = next((c for c in cands if c in present), None)
        chosen.append((header, col))
        if header != CAT_HEADER and col is None and cands:
            missing.append(f"{header} (wanted {cands[0]})")
    return chosen, missing


def refine(raw_headers, raw_rows, cat_lut):
    chosen, missing = resolve_sources(raw_headers)
    if missing:
        print("*** WARNING: raw columns not found, left blank: "
              + "; ".join(missing))
    out_headers = [h for h, _ in chosen]
    out_rows = []
    hits = 0
    for r in raw_rows:
        row = {}
        for header, col in chosen:
            row[header] = "" if col is None else r.get(col, "")
        if cat_lut:
            supplier = nkey(row.get(MATCH_SOURCE_HEADER, ""))
            cat = cat_lut.get(supplier, "")
            row[CAT_HEADER] = cat
            if cat:
                hits += 1
        out_rows.append(row)
    if cat_lut:
        print(f"  Cat filled for {hits}/{len(out_rows)} rows "
              f"({len(out_rows) - hits} had no supplier match)")
    return out_headers, out_rows


def write_csv(path, headers, rows):
    target = path
    for attempt in range(10):
        try:
            with open(target, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
                w.writeheader()
                w.writerows(rows)
            break
        except PermissionError:
            base, ext = Path(path).stem, Path(path).suffix
            target = str(Path(path).with_name(f"{base}_{attempt + 1}{ext}"))
            print(f"    (file locked, trying {target})")
    else:
        sys.exit(f"Could not write {path} - close it in Excel and retry.")
    print(f"\nSaved: {target}  ({len(rows)} rows, {len(headers)} columns)")


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(
        description="Refine a raw OneTrust export to the exact 19 GUI columns.")
    ap.add_argument("--in", dest="inp", required=True, help="Raw export (.csv).")
    ap.add_argument("--out", default="refined.csv", help="Output CSV path.")
    ap.add_argument("--cat-file", default="Supplier_Category_List.xlsx",
                    help="Supplier->category lookup (.xlsx/.csv). Skipped if absent.")
    ap.add_argument("--cat-key", help="Supplier-name column in the lookup sheet.")
    ap.add_argument("--cat-value", help="Category column in the lookup sheet.")
    args = ap.parse_args()

    raw_headers, raw_rows = load_table(args.inp)
    if not raw_rows:
        sys.exit("Input has no data rows.")

    cat_lut = {}
    if Path(args.cat_file).exists():
        cat_lut = build_cat_lookup(args.cat_file, args.cat_key, args.cat_value)
    else:
        print(f"  (no {args.cat_file} found -> 'Cat' left blank)")

    out_headers, out_rows = refine(raw_headers, raw_rows, cat_lut)
    write_csv(args.out, out_headers, out_rows)


if __name__ == "__main__":
    main()

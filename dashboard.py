#!/usr/bin/env python3
"""Local dashboard server for OneTrust zone risk insights."""
import argparse
import datetime
import json
import subprocess
import sys
import threading
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from zone_reports import stats
from zone_reports import narrative
import os
import tempfile
from config_store import load_config, save_config, is_configured, next_run_due, host_hint
from onetrust_view_export import get_token

HERE = Path(__file__).resolve().parent


def zone_to_dict(zs, narrative_lines):
    return {
        "organization": zs.organization,
        "total": zs.total,
        "by_cat": zs.by_cat,
        "grand_by_cat": zs.grand_by_cat,
        "by_domain": zs.by_domain,
        "crosstab": zs.crosstab,
        "domain_pct": zs.domain_pct,
        "top_cats": zs.top_cats,
        "treated": zs.treated,
        "treated_pct": zs.treated_pct,
        "by_stage": zs.by_stage,
        "narrative": narrative_lines,
    }


def build_payload(rows, view=None, generated_at=None, source=None):
    generated_at = generated_at or datetime.datetime.now()
    date_str = narrative.format_date_ordinal(generated_at.date())
    zones = {}
    for code, org in stats.ZONES.items():
        zs = stats.compute_zone_stats(rows, org)
        lines = narrative.build_narrative(zs, date_str)
        zones[code] = zone_to_dict(zs, lines)
    payload = {
        "generated_at": generated_at.replace(microsecond=0).isoformat(),
        "view": view,
        "cat_order": list(stats.CAT_ORDER),
        "zones": zones,
    }
    if source is not None:
        payload["source"] = source
    return payload


RAW = str(HERE / "raw_export.csv")
DATA = str(HERE / "refined.csv")
CAT_FILE = str(HERE / "Supplier_Category_List.xlsx")
VIEWS_FILE = HERE / "views.json"


def export_cmd(view, hostname, client_id, client_secret, raw, py=sys.executable):
    return [py, str(HERE / "onetrust_view_export.py"),
            "--view", view, "--all-columns", "--out", raw,
            "--hostname", hostname, "--client-id", client_id,
            "--client-secret", client_secret]


def refine_cmd(raw, data, cat_file, py=sys.executable):
    return [py, str(HERE / "refine_columns.py"),
            "--in", raw, "--out", data, "--cat-file", cat_file]


def _run_step(cmd, runner):
    r = runner(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        tail = (r.stderr or r.stdout or "").strip().splitlines()
        msg = "\n".join(tail[-5:]) if tail else f"exit {r.returncode}"
        raise RuntimeError(msg)


CONFIG_FILE = HERE / "config.json"
VIEW = "TPRM Global View"

_job_lock = threading.Lock()
JOB = {"state": "idle", "last_run": None, "rows": None, "message": "", "source": None}


def reset_job():
    global JOB
    JOB = {"state": "idle", "last_run": None, "rows": None, "message": "", "source": None}


def do_fetch(cfg, *, runner=subprocess.run, py=sys.executable):
    if not is_configured(cfg):
        raise ValueError("Not configured")
    _run_step(export_cmd(VIEW, cfg["hostname"], cfg["client_id"],
                         cfg["client_secret"], RAW, py), runner)
    _run_step(refine_cmd(RAW, DATA, CAT_FILE, py), runner)
    return len(stats.load_rows(DATA))


def run_fetch_job(*, config_loader=None, runner=subprocess.run,
                  py=sys.executable, now_fn=None):
    config_loader = config_loader or (lambda: load_config(str(CONFIG_FILE)))
    now_fn = now_fn or datetime.datetime.now
    try:
        n = do_fetch(config_loader(), runner=runner, py=py)
        JOB.update(state="done", rows=n, last_run=now_fn(), source="api", message="")
    except Exception as e:
        JOB.update(state="error", message=str(e))


def start_fetch_async():
    with _job_lock:
        if JOB["state"] == "running":
            return {"started": False, "running": True}
        JOB["state"] = "running"
        JOB["message"] = ""
    t = threading.Thread(target=run_fetch_job, daemon=True)
    t.start()
    return {"started": True, "running": False}


def scheduler_tick(now, *, config_loader=None):
    cfg = (config_loader or (lambda: load_config(str(CONFIG_FILE))))()
    if not is_configured(cfg):
        return False
    if next_run_due(cfg.get("schedule", {}), JOB["last_run"], now) and JOB["state"] != "running":
        return start_fetch_async().get("started", False)
    return False


def _scheduler_loop(stop_event):
    while not stop_event.wait(30):
        try:
            scheduler_tick(datetime.datetime.now())
        except Exception:
            pass  # never let the daemon die on a transient error


def test_connection(hostname, client_id, client_secret):
    try:
        get_token(hostname, client_id, client_secret)
        return {"ok": True, "message": "Connection OK"}
    except SystemExit as e:
        return {"ok": False, "message": str(e)}
    except Exception as e:
        return {"ok": False, "message": f"{type(e).__name__}: {e}"}


def load_cached():
    if not Path(DATA).exists():
        return None
    rows = stats.load_rows(DATA)
    return build_payload(rows, source="cache")


def list_views():
    if not VIEWS_FILE.exists():
        return []
    try:
        data = json.loads(VIEWS_FILE.read_text(encoding="utf-8"))
        return [v.get("name", "") for v in data.get("views", []) if v.get("name")]
    except (json.JSONDecodeError, OSError):
        return []


def strict_mask(cfg):
    return {"configured": is_configured(cfg),
            "host_hint": host_hint(cfg.get("hostname")),
            "has_secret": bool((cfg.get("client_secret") or "").strip()),
            "schedule": dict(cfg.get("schedule") or {})}


def _count_views():
    try:
        data = json.loads(Path(VIEWS_FILE).read_text(encoding="utf-8"))
        return len(data.get("views", []))
    except Exception:
        return None


def _count_suppliers():
    try:
        from openpyxl import load_workbook
        wb = load_workbook(CAT_FILE, read_only=True)
        n = max(0, (wb.active.max_row or 1) - 1)
        wb.close()
        return n
    except Exception:
        return None


def files_status():
    cat_present = Path(CAT_FILE).exists()
    views_present = Path(VIEWS_FILE).exists()
    return {"catfile": {"present": cat_present,
                        "suppliers": _count_suppliers() if cat_present else None},
            "views": {"present": views_present,
                      "view_count": _count_views() if views_present else None}}


def _atomic_write(target, data):
    d = os.path.dirname(str(target)) or "."
    fd, tmp = tempfile.mkstemp(dir=d, suffix=".tmp")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        os.replace(tmp, str(target))
    except Exception:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


INDEX = HERE / "index.html"


def _json_resp(obj, status=200):
    return status, "application/json; charset=utf-8", json.dumps(obj).encode("utf-8")


def _extra_header(path, status):
    """Return (header_name, header_value) tuple for Content-Disposition, or None."""
    if status != 200:
        return None
    if path == "/api/ppt":
        return ("Content-Disposition", 'attachment; filename="risk-insights.pptx"')
    if path == "/api/catfile":
        return ("Content-Disposition", 'attachment; filename="Supplier_Category_List.xlsx"')
    return None


def _job_status():
    j = JOB
    return {"state": j["state"],
            "last_run": j["last_run"].isoformat() if j["last_run"] else None,
            "rows": j["rows"], "message": j["message"], "source": j["source"]}


def _handle_ppt():
    payload = load_cached()
    if not payload:
        return _json_resp({"error": "No data — refresh first"}, 400)
    try:
        import ppt_export
        data = ppt_export.deck_to_bytes(payload)
    except Exception as e:
        return _json_resp({"error": f"PPT build failed: {e}"}, 500)
    ctype = ("application/vnd.openxmlformats-officedocument"
             ".presentationml.presentation")
    return 200, ctype, data


def handle_get(path):
    if path == "/" or path == "/index.html":
        try:
            return 200, "text/html; charset=utf-8", INDEX.read_bytes()
        except FileNotFoundError:
            return _json_resp({"error": "index.html not found"}, 500)
    if path == "/api/config":
        return _json_resp(strict_mask(load_config(str(CONFIG_FILE))))
    if path == "/api/files":
        return _json_resp(files_status())
    if path == "/api/catfile":
        if not Path(CAT_FILE).exists():
            return _json_resp({"error": "No supplier file uploaded"}, 404)
        ctype = ("application/vnd.openxmlformats-officedocument"
                 ".spreadsheetml.sheet")
        return 200, ctype, Path(CAT_FILE).read_bytes()
    if path == "/api/status":
        return _json_resp(_job_status())
    if path == "/api/views":
        return _json_resp({"views": list_views()})
    if path == "/api/data":
        payload = load_cached()
        return _json_resp(payload if payload else {"empty": True})
    if path == "/api/ppt":
        return _handle_ppt()          # defined in Task 5
    return _json_resp({"error": "not found"}, 404)


def handle_post(path, body_bytes):
    try:
        if path == "/api/fetch":
            r = start_fetch_async()
            return _json_resp(r, 202)
        if path == "/api/config":
            body = json.loads(body_bytes or b"{}")
            merged = save_config(str(CONFIG_FILE), body, load_config(str(CONFIG_FILE)))
            return _json_resp(strict_mask(merged))
        if path == "/api/upload/catfile":
            if not body_bytes or len(body_bytes) > 10 * 1024 * 1024:
                return _json_resp({"error": "File empty or over 10 MB"}, 400)
            if not body_bytes.startswith(b"PK"):
                return _json_resp({"error": "Not a valid .xlsx file"}, 400)
            _atomic_write(CAT_FILE, body_bytes)
            return _json_resp(files_status())
        if path == "/api/upload/views":
            if not body_bytes or len(body_bytes) > 2 * 1024 * 1024:
                return _json_resp({"error": "File empty or over 2 MB"}, 400)
            try:
                parsed = json.loads(body_bytes)
                if not isinstance(parsed.get("views"), list):
                    raise ValueError("missing views list")
            except Exception:
                return _json_resp({"error": "Not a valid views.json (needs a 'views' list)"}, 400)
            _atomic_write(VIEWS_FILE, body_bytes)
            return _json_resp(files_status())
        if path == "/api/config/test":
            body = json.loads(body_bytes or b"{}")
            existing = load_config(str(CONFIG_FILE))
            host = body.get("hostname") or existing["hostname"]
            cid = body.get("client_id") or existing["client_id"]
            sec = body.get("client_secret") or existing["client_secret"]
            return _json_resp(test_connection(host, cid, sec))
        return _json_resp({"error": "not found"}, 404)
    except Exception as e:
        return _json_resp({"error": str(e)}, 500)


def make_handler():
    class Handler(BaseHTTPRequestHandler):
        def _send(self, status, ctype, body, extra=None):
            self.send_response(status)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(body)))
            if extra:
                self.send_header(*extra)
            self.end_headers()
            self.wfile.write(body)

        def do_GET(self):
            status, ctype, body = handle_get(self.path)
            extra = _extra_header(self.path, status)
            self._send(status, ctype, body, extra)

        def do_POST(self):
            length = max(0, int(self.headers.get("Content-Length", 0)))
            body = self.rfile.read(length) if length else b""
            self._send(*handle_post(self.path, body))

        def log_message(self, *a):
            pass  # silence; never log request bodies (may hold secrets)
    return Handler


def main(argv=None):
    ap = argparse.ArgumentParser(description="OneTrust risk insights dashboard.")
    ap.add_argument("--port", type=int, default=8000)
    ap.add_argument("--no-browser", action="store_true")
    args = ap.parse_args(argv)
    url = f"http://127.0.0.1:{args.port}/"
    # Threading so a slow request (e.g. Test Connection reaching OneTrust)
    # never blocks concurrent requests like Save. daemon_threads: no hang on exit.
    server = ThreadingHTTPServer(("127.0.0.1", args.port), make_handler())
    server.daemon_threads = True
    print(f"Dashboard at {url}  (Ctrl-C to stop)")
    stop_event = threading.Event()
    threading.Thread(target=_scheduler_loop, args=(stop_event,), daemon=True).start()
    if not args.no_browser:
        webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")
        stop_event.set()
        server.server_close()


if __name__ == "__main__":
    main()
